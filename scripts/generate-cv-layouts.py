# -*- coding: utf-8 -*-
"""生成 4 套简历排版方案的 Word 模板(占位符版)。

用法:
    python scripts/generate-cv-layouts.py

输出到 templates/layouts/:
    cv-layout-a-ats.docx / cv-layout-b-tech.docx /
    cv-layout-c-compact.docx / cv-layout-d-interview.docx

设计规范: knowledge/resume-templates/layout-schemes.md
内容骨架: templates/cv-template.md(纯占位符,不含真实个人信息)
PDF 转换: soffice --headless --convert-to pdf --outdir templates/layouts <docx>
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "templates", "layouts")

EAST_FONT = "微软雅黑"
ASCII_FONT = "Arial"
DARK_BLUE = "1F3864"
GRAY_TAG = "595959"
GRAY_NOTE = "808080"
LIGHT_GRAY_BG = "F2F2F2"

# ---------------------------------------------------------------- 内容占位符

NAME = "{{name}}"
INTENT = "求职意向:{{target_primary}} | {{target_cities}} | 期望薪资 {{salary_min}}-{{salary_max}}k"
CONTACT = "[手机号] | [邮箱]"
EDU_LINE = "{{school}} | {{major}} | {{degree}} | {{graduation}}毕业"
SKILLS = [
    ("熟练", "{{skills_proficient}}"),
    ("掌握", "{{skills_familiar}}"),
    ("学习中", "{{skills_learning}}"),
]
# (名称, 年份, 技术标签行(方案B/D用), [bullet...])
BULLET_DUTY = "[一句话职责]"
BULLET_METRIC = "[量化成果,例如:在线率从 X% 提升至 Y%]"
BULLET_TECH = "[关键技术/排查过程]"
PROJECTS = [
    ("{{project_1_name}}", "{{project_1_year}}", "[技术栈标签,例如:ESP32 · MQTT · Docker · AWS IoT Core]",
     [BULLET_DUTY, BULLET_METRIC, BULLET_TECH]),
    ("{{project_2_name}}", "{{project_2_year}}", "[技术栈标签]",
     [BULLET_DUTY, BULLET_METRIC, BULLET_TECH]),
    ("{{project_3_name}}", "{{project_3_year}}", "[技术栈标签]",
     [BULLET_DUTY, BULLET_METRIC, BULLET_TECH]),
    ("{{project_4_name}}", "{{project_4_year}}", "[技术栈标签]",
     [BULLET_DUTY, BULLET_METRIC, BULLET_TECH]),
]
CERTS = ["{{cert_1}}", "[其他证书]"]

NOTES = {
    "a": "方案 A · 经典单栏 ATS 版 —— 适用于官网/招聘平台投递;纯黑白、单栏、无表格无文本框,ATS 解析最稳。",
    "b": "方案 B · 技术标签项目版 —— 适用于 IoT/AI Infra/云平台等技术岗投递;技能前置,标签行突出 JD 关键词。",
    "c": "方案 C · 紧凑海投一页版 —— 仅限海投/平台附件/快速沟通等明确一页要求的场景;正式投递请用方案 A/B。",
    "d": "方案 D · 面谈携带人读版 —— 仅限面试现场/内推直发,禁止上传任何招聘平台(双栏+色块会导致 ATS 解析错乱)。",
}

# ---------------------------------------------------------------- 基础 helper


def set_run(run, size=10.5, bold=False, color=None):
    """统一设置中西文字体:ASCII 用 Arial,中文用微软雅黑。"""
    run.font.name = ASCII_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ASCII_FONT)
    rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    rFonts.set(qn("w:eastAsia"), EAST_FONT)


def para(doc_or_cell, text="", size=10.5, bold=False, color=None,
         before=0.0, after=2.0, line=None):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if text:
        set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_bottom_border(p, sz=6, color="000000"):
    """段落底边框,实现模块标题的通栏细分隔线(0.75pt = sz 6)。"""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def clear_cell(cell):
    """移除单元格自带的空首段落,避免分栏顶部多出空行。"""
    for p in list(cell.paragraphs):
        if not p.text:
            p._p.getparent().remove(p._p)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_col_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def new_doc(margin_tb, margin_lr):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.top_margin = sec.bottom_margin = Cm(margin_tb)
    sec.left_margin = sec.right_margin = Cm(margin_lr)
    return doc


def add_project(doc, name, year, tag, bullets, body_size, line,
                tab_at_cm, tag_line=False, metric_first=False,
                title_before=4, item_after=1):
    """项目标题(年份右对齐同行) + 可选技术标签行 + bullets。"""
    p = para(doc, before=title_before, after=item_after, line=line)
    set_run(p.add_run(name), size=11 if body_size >= 10.5 else 10.5, bold=True)
    p.add_run("\t")
    set_run(p.add_run(year), size=body_size)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(tab_at_cm), WD_TAB_ALIGNMENT.RIGHT)
    if tag_line and tag:
        para(doc, tag, size=10, color=GRAY_TAG, after=item_after, line=line)
    items = list(bullets)
    if metric_first and BULLET_METRIC in items:
        items.remove(BULLET_METRIC)
        items.insert(0, BULLET_METRIC)
    for b in items:
        para(doc, "- " + b, size=body_size, after=item_after, line=line)


def add_footer_note(doc, text, line=None, before=8):
    """页脚位置的小字说明(普通段落,非真实页脚)。"""
    para(doc, text, size=8, color=GRAY_NOTE, before=before, after=0, line=line)


def section_title(doc, text, size=12, border=True, color=None, before=8, after=4):
    p = para(doc, text, size=size, bold=True, color=color, before=before, after=after)
    if border:
        add_bottom_border(p, sz=6, color="000000")
    return p


def add_header(doc, name_size, contact=True):
    para(doc, NAME, size=name_size, bold=True, after=2)
    para(doc, INTENT, size=10.5, after=1)
    if contact:
        para(doc, CONTACT, size=10.5, after=1)


def add_skills(doc, body_size, line, cell=None):
    target = cell if cell is not None else doc
    for label, value in SKILLS:
        p = para(target, after=1, line=line)
        set_run(p.add_run(label + ":"), size=body_size, bold=True)
        set_run(p.add_run(value), size=body_size)


def add_edu(doc, body_size, line, cell=None):
    para(cell if cell is not None else doc, EDU_LINE, size=body_size, after=1, line=line)


def add_certs(doc, body_size, line, cell=None):
    target = cell if cell is not None else doc
    for c in CERTS:
        para(target, "- " + c, size=body_size, after=1, line=line)


# ---------------------------------------------------------------- 方案 A


def build_a():
    doc = new_doc(2.0, 2.2)
    tab_at = 21.0 - 2.2 * 2  # 通栏右边距处
    add_header(doc, 20)
    section_title(doc, "教育背景")
    add_edu(doc, 10.5, 19)
    section_title(doc, "技能")
    add_skills(doc, 10.5, 19)
    section_title(doc, "项目经历")
    for name, year, tag, bullets in PROJECTS:
        add_project(doc, name, year, tag, bullets, 10.5, 19, tab_at)
    section_title(doc, "证书与竞赛")
    add_certs(doc, 10.5, 19)
    add_footer_note(doc, NOTES["a"])
    return doc


# ---------------------------------------------------------------- 方案 B


def build_b():
    doc = new_doc(2.0, 2.2)
    tab_at = 21.0 - 2.2 * 2
    add_header(doc, 20)
    # 模块顺序:意向 → 技能 → 项目经历 → 教育 → 证书(技能前置)
    # 段距收紧(字号/行距仍遵守规范下限)以保证 4 项目 + 标签行在一页内
    section_title(doc, "技能", color=DARK_BLUE, before=5, after=3)
    add_skills(doc, 10.5, 18)
    section_title(doc, "项目经历", color=DARK_BLUE, before=5, after=3)
    for name, year, tag, bullets in PROJECTS:
        add_project(doc, name, year, tag, bullets, 10.5, 18, tab_at,
                    tag_line=True, metric_first=True, title_before=2, item_after=0)
    section_title(doc, "教育背景", color=DARK_BLUE, before=5, after=3)
    add_edu(doc, 10.5, 18)
    section_title(doc, "证书与竞赛", color=DARK_BLUE, before=5, after=3)
    add_certs(doc, 10.5, 18)
    add_footer_note(doc, NOTES["b"], before=4)
    return doc


# ---------------------------------------------------------------- 方案 C


def build_c():
    doc = new_doc(1.5, 1.8)
    tab_at = 21.0 - 1.8 * 2
    para(doc, NAME, size=18, bold=True, after=1)
    para(doc, INTENT, size=10, after=0, line=16)
    para(doc, CONTACT, size=10, after=1, line=16)
    # 模块标题:11pt 加粗、无分隔线、段前 6pt
    section_title(doc, "技能", size=11, border=False, before=6, after=2)
    add_skills(doc, 10, 16)
    section_title(doc, "项目经历", size=11, border=False, before=6, after=2)
    for name, year, tag, bullets in PROJECTS[:3]:  # 只保留 3 个最强项目
        add_project(doc, name, year, tag, bullets, 10, 16, tab_at)
    section_title(doc, "教育与证书", size=11, border=False, before=6, after=2)
    add_edu(doc, 10, 16)
    para(doc, "{{cert_1}} | [其他证书]", size=10, after=1, line=16)
    add_footer_note(doc, NOTES["c"], line=16)
    return doc


# ---------------------------------------------------------------- 方案 D


def build_d():
    doc = new_doc(1.5, 1.5)
    content_w = 21.0 - 1.5 * 2  # 18cm
    left_w = round(content_w * 0.30, 1)   # 5.4cm
    right_w = round(content_w - left_w, 1)  # 12.6cm

    para(doc, NAME, size=20, bold=True, color=DARK_BLUE, after=1)
    para(doc, INTENT, size=10.5, after=4)

    table = doc.add_table(rows=1, cols=2)
    set_col_widths(table, [left_w, right_w])
    left, right = table.rows[0].cells
    clear_cell(left)
    clear_cell(right)
    shade_cell(left, LIGHT_GRAY_BG)

    # 左窄栏:联系方式 / 技能 / 证书
    para(left, "联系方式", size=11, bold=True, color=DARK_BLUE, before=2, after=2)
    para(left, "[手机号]", size=10, after=1)
    para(left, "[邮箱]", size=10, after=1)
    para(left, "技能", size=11, bold=True, color=DARK_BLUE, before=8, after=2)
    add_skills(doc, 10, None, cell=left)
    para(left, "证书与竞赛", size=11, bold=True, color=DARK_BLUE, before=8, after=2)
    add_certs(doc, 10, None, cell=left)

    # 右主栏:项目经历 / 教育背景
    tab_at = right_w - 0.4  # 扣除单元格默认左右边距
    para(right, "项目经历", size=12, bold=True, color=DARK_BLUE, before=2, after=3)
    for name, year, tag, bullets in PROJECTS:
        add_project(right, name, year, tag, bullets, 10, None, tab_at, tag_line=True)
    para(right, "教育背景", size=12, bold=True, color=DARK_BLUE, before=8, after=3)
    add_edu(doc, 10, None, cell=right)

    add_footer_note(doc, NOTES["d"])
    return doc


# ---------------------------------------------------------------- main


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    builders = {
        "cv-layout-a-ats.docx": build_a,
        "cv-layout-b-tech.docx": build_b,
        "cv-layout-c-compact.docx": build_c,
        "cv-layout-d-interview.docx": build_d,
    }
    for filename, builder in builders.items():
        path = os.path.join(OUT_DIR, filename)
        builder().save(path)
        print("saved:", os.path.normpath(path))


if __name__ == "__main__":
    main()
