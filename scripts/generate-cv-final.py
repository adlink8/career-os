# -*- coding: utf-8 -*-
"""解析三份定稿简历 md,生成正式 docx + PDF,以及项目 Portfolio(docx/pdf/html)。

用法:
    python scripts/generate-cv-final.py

输入(只读,不得改写文案):
    data/cv/cv-ai-infra-v2-draft.md -> cv-ai-infra.docx/.pdf
    data/cv/cv-iot-v2-draft.md      -> cv-iot.docx/.pdf
    data/cv/cv-ops-v2-draft.md      -> cv-ops.docx/.pdf

输出到 data/cv/final/(已 gitignore):
    cv-*.docx/.pdf、portfolio.docx/.pdf、portfolio.html

排版规范: knowledge/resume-templates/layout-schemes.md 方案 B
(A4,上下 2.0cm/左右 2.2cm;姓名 20pt 加粗;模块标题 12pt 加粗深蓝 #1F3864 +
通栏细分隔线 sz=6;正文 10.5pt 固定行距 18pt;项目标题 11pt 加粗 +
灰色技术标签行 10pt #595959;模块顺序:意向→技能→项目→实习→教育→证书)
PDF 转换: LibreOffice soffice --headless --convert-to pdf
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CV_DIR = os.path.join(ROOT, "data", "cv")
OUT_DIR = os.path.join(CV_DIR, "final")
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

EAST_FONT = "微软雅黑"
ASCII_FONT = "Arial"
DARK_BLUE = "1F3864"
GRAY_TAG = "595959"

CV_FILES = {
    "cv-ai-infra": "cv-ai-infra-v2-draft.md",
    "cv-iot": "cv-iot-v2-draft.md",
    "cv-ops": "cv-ops-v2-draft.md",
}

# ---------------------------------------------------------------- md 解析


def parse_cv(path):
    """把简历 md 解析为结构化 dict。占位符([手机号] 等)原样保留。"""
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    data = {
        "name": "", "intent": "", "contact": "",
        "skills": [],        # [(label, value)]
        "projects": [],      # [{title, tag, bullets}]
        "internship": {"title_bold": "", "title_rest": "", "bullets": []},
        "education": [],     # [line]
        "certs_title": "证书",
        "certs": [],
    }
    section = None
    project = None
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            data["name"] = re.sub(r"\(.*$", "", line[2:]).strip()
            continue
        if line.startswith("## "):
            section = line[3:].strip()
            project = None
            if section.startswith("证书"):
                data["certs_title"] = section
            continue
        if line.startswith("**求职意向**"):
            data["intent"] = line.split(":", 1)[1].strip()
            continue
        if line.startswith("**联系方式**"):
            data["contact"] = line.split(":", 1)[1].strip()
            continue

        if section == "技能":
            m = re.match(r"^- \*\*(.+?)\*\*:(.*)$", line)
            if m:
                data["skills"].append((m.group(1).strip(), m.group(2).strip()))
        elif section == "项目经历":
            if line.startswith("### "):
                project = {"title": line[4:].strip(), "tag": "", "bullets": []}
                data["projects"].append(project)
            elif project is not None and line.startswith("`") and line.endswith("`"):
                project["tag"] = line.strip("`").strip()
            elif project is not None and line.startswith("- "):
                project["bullets"].append(line[2:].strip())
        elif section == "实习经历":
            if line.startswith("**"):
                parts = line.split("**")
                data["internship"]["title_bold"] = parts[1] if len(parts) > 1 else line
                data["internship"]["title_rest"] = parts[2].strip() if len(parts) > 2 else ""
            elif line.startswith("- "):
                data["internship"]["bullets"].append(line[2:].strip())
        elif section == "教育背景":
            data["education"].append(line)
        elif section and section.startswith("证书"):
            if line.startswith("- "):
                data["certs"].append(line[2:].strip())
    return data


# ---------------------------------------------------------------- docx helper
# (函数风格复用 scripts/generate-cv-layouts.py)


def set_run(run, size=10.5, bold=False, color=None):
    """统一设置中西文字体:ASCII 用 Arial,中文用微软雅黑。"""
    run.font.name = ASCII_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ASCII_FONT)
    rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    rFonts.set(qn("w:eastAsia"), EAST_FONT)


def para(doc, text="", size=10.5, bold=False, color=None,
         before=0.0, after=0.0, line=18):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if text:
        set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_bottom_border(p, sz=6, color="000000"):
    """段落底边框,实现模块标题的通栏细分隔线(0.75pt = sz 6)。"""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_title(doc, text, before=2, after=1):
    p = para(doc, text, size=12, bold=True, color=DARK_BLUE,
             before=before, after=after, line=15)
    add_bottom_border(p, sz=6, color="000000")
    return p


def new_doc(margin_tb=2.0, margin_lr=2.2):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.top_margin = sec.bottom_margin = Cm(margin_tb)
    sec.left_margin = sec.right_margin = Cm(margin_lr)
    return doc


def metric_first(bullets):
    """量化成果(含 % 的 bullet)前置,其余保持原顺序。"""
    return sorted(bullets, key=lambda b: 0 if "%" in b else 1)


# ---------------------------------------------------------------- CV 生成


def build_cv(data, out_path):
    doc = new_doc()
    # 头部:姓名 20pt 加粗;意向/联系 10.5pt
    para(doc, data["name"], size=20, bold=True, after=1, line=24)
    p = para(doc, after=0)
    set_run(p.add_run("求职意向:"), size=10.5, bold=True)
    set_run(p.add_run(data["intent"]), size=10.5)
    p = para(doc, after=0)
    set_run(p.add_run("联系方式:"), size=10.5, bold=True)
    set_run(p.add_run(data["contact"]), size=10.5)

    # 技能(前置)
    section_title(doc, "技能")
    for label, value in data["skills"]:
        p = para(doc, after=0)
        set_run(p.add_run(label + ":"), size=10.5, bold=True)
        set_run(p.add_run(value), size=10.5)

    # 项目经历:标题 11pt 加粗 + 灰色技术标签行 + bullets
    section_title(doc, "项目经历")
    for proj in data["projects"]:
        para(doc, proj["title"], size=11, bold=True, before=1, after=0)
        if proj["tag"]:
            para(doc, proj["tag"], size=10, color=GRAY_TAG, after=0, line=16)
        for b in metric_first(proj["bullets"]):
            para(doc, "- " + b, size=10.5, after=0)

    # 实习经历
    section_title(doc, "实习经历")
    ins = data["internship"]
    if ins["title_bold"]:
        p = para(doc, after=0)
        set_run(p.add_run(ins["title_bold"]), size=10.5, bold=True)
        if ins["title_rest"]:
            set_run(p.add_run(" " + ins["title_rest"]), size=10.5)
    for b in ins["bullets"]:
        para(doc, "- " + b, size=10.5, after=0)

    # 教育背景
    section_title(doc, "教育背景")
    for line in data["education"]:
        para(doc, line, size=10.5, after=0)

    # 证书
    section_title(doc, data["certs_title"])
    for c in data["certs"]:
        para(doc, "- " + c, size=10.5, after=0)

    doc.save(out_path)
    print("saved:", os.path.normpath(out_path))


# ---------------------------------------------------------------- Portfolio 内容
# 事实与数字严格取自三份草稿,不得新造。

PORTFOLIO = {
    "name": "[姓名]",
    "title": "项目 Portfolio",
    "contact": "GitHub: github.com/your-username",
    "projects": [
        {
            "name": "NovelMind — 中文长文本 RAG 平台(独立项目)",
            "tag": "Python / FastAPI / Next.js / PostgreSQL / ChromaDB / Ollama",
            "oneline": "独立设计的长文本语义检索与问答 RAG 平台,以小说阅读为验证场景,覆盖前后端与数据层。",
            "arch": [
                "RAG 管道全链路:切块 → embedding → 向量库 → 本地大模型生成,覆盖前后端与数据层",
                "质量评估闭环:固定测试集 + Recall@K / MRR / faithfulness 指标 + 错误案例分析",
                "GSD 阶段化管理迭代,pytest 自动化回归保障每次变更可验证",
            ],
            "results": [
                "建立 Recall@K / MRR / faithfulness 三指标评估闭环,切块与检索策略迭代均由指标驱动,而非凭感觉调参",
                "独立交付覆盖前后端与数据层的完整平台,配套 pytest 测试套件,关键变更跑回归验证",
            ],
            "debug": "以小说阅读为验证场景:长文本的人物关系、时间线、伏笔检索最能暴露切块质量问题,据此定位并迭代切块与检索策略。",
            "link": "https://github.com/adlink8/novel-mind",
        },
        {
            "name": "T5AI 设备数据监控终端(固件 + PC 桥接,开源)",
            "tag": "C / TuyaOpen SDK / LVGL / MQTT / Mosquitto",
            "oneline": "涂鸦 T5AI 开发板上的设备数据监控终端:LVGL 固件 + PC 桥接服务,代码已开源。",
            "arch": [
                "双通道架构:PC bridge_server → Mosquitto MQTT Broker → 开发板为主通道,HTTP 轮询为回退",
                "LVGL 应用(480×320 LCD)实时显示 PC 端推送的数据",
                "基于 TuyaOpen SDK(WSL)完成固件构建、烧录与回归测试",
            ],
            "results": [
                "MQTT 主通道 + HTTP 回退双通道架构落地,断连场景下数据仍可送达",
                "代码开源:GitHub adlink8/t5ai-codex-quota",
            ],
            "debug": "在 WSL 环境完成固件构建、烧录与回归测试,解决中文字形缺失、串口烧录卡死等问题。",
            "link": "https://github.com/adlink8/t5ai-codex-quota",
        },
        {
            "name": "个人数据基础设施(多源数据统合 + 知识检索服务)",
            "tag": "Python / SQLite / Chroma / MCP / REST / CLI",
            "oneline": "多源数据统合与知识检索服务,长期在线运行并被个人知识管理与求职系统持续调用。",
            "arch": [
                "多源数据统合管道:AI 对话/行为事件 → 清洗 → 结构化入库",
                "对外提供 CLI / REST / MCP 三种检索接口",
                "索引版本化与快照机制:构建 → 校验 → 激活 → 可回滚",
            ],
            "results": [
                "沉淀 11,000+ 事件、44,000+ 实体、40,000+ 知识单元索引",
                "长期在线运行并被个人知识管理与求职系统持续调用,具备快照快速恢复能力",
            ],
            "debug": "以索引版本化与快照机制(构建 → 校验 → 激活 → 可回滚)应对索引损坏与误变更,按生产标准维护服务稳定性。",
            "link": "本地私有仓库,可现场演示",
        },
    ],
    "ai_collab": [
        ("设计归我", "架构/选型/验收标准我定,AI 加速实现不替代判断"),
        ("阶段化协作", "GSD 规划→执行→验证→审计,每阶段独立验收"),
        ("先验证再宣称", "测试要过、服务要跑、数字要有出处"),
        ("AI 是工具链不是拐杖", "MCP/Agent/RAG 是构建方式,上述项目即证据"),
    ],
}


def build_portfolio_docx(out_path):
    doc = new_doc()
    p = para(doc, PORTFOLIO["name"] + " · " + PORTFOLIO["title"],
             size=16, bold=True, color=DARK_BLUE, after=1, line=None)
    para(doc, PORTFOLIO["contact"], size=10.5, color=GRAY_TAG, after=2)

    for proj in PORTFOLIO["projects"]:
        section_title(doc, proj["name"], before=6, after=2)
        para(doc, proj["oneline"], size=10.5, after=1)
        para(doc, proj["tag"], size=10, color=GRAY_TAG, after=1)
        p = para(doc, after=0)
        set_run(p.add_run("架构要点"), size=10.5, bold=True)
        for b in proj["arch"]:
            para(doc, "- " + b, size=10.5, after=0)
        p = para(doc, after=0, before=2)
        set_run(p.add_run("关键成果"), size=10.5, bold=True)
        for b in proj["results"]:
            para(doc, "- " + b, size=10.5, after=0)
        p = para(doc, after=0, before=2)
        set_run(p.add_run("排障故事"), size=10.5, bold=True)
        para(doc, "- " + proj["debug"], size=10.5, after=0)
        p = para(doc, after=0, before=2)
        set_run(p.add_run("链接:"), size=10.5, bold=True)
        set_run(p.add_run(proj["link"]), size=10.5)

    section_title(doc, "我与 AI 的协作模式", before=6, after=2)
    for head, body in PORTFOLIO["ai_collab"]:
        p = para(doc, after=1)
        set_run(p.add_run("- " + head + ":"), size=10.5, bold=True)
        set_run(p.add_run(body), size=10.5)

    doc.save(out_path)
    print("saved:", os.path.normpath(out_path))


# ---------------------------------------------------------------- Portfolio HTML


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_portfolio_html(out_path):
    cards = []
    for proj in PORTFOLIO["projects"]:
        arch = "\n".join("        <li>%s</li>" % esc(b) for b in proj["arch"])
        results = "\n".join("        <li>%s</li>" % esc(b) for b in proj["results"])
        cards.append("""    <section class="card">
      <h2>%s</h2>
      <p class="oneline">%s</p>
      <p class="tag">%s</p>
      <h3>架构要点</h3>
      <ul>
%s
      </ul>
      <h3>关键成果</h3>
      <ul>
%s
      </ul>
      <h3>排障故事</h3>
      <p class="debug">%s</p>
      <p class="link">链接:%s</p>
    </section>""" % (
            esc(proj["name"]), esc(proj["oneline"]), esc(proj["tag"]),
            arch, results, esc(proj["debug"]), esc(proj["link"])))

    collab = "\n".join(
        "      <li><strong>%s</strong> — %s</li>" % (esc(h), esc(b))
        for h, b in PORTFOLIO["ai_collab"])

    html = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>%s · %s</title>
<style>
  * { margin: 0; padding: 0; box-sizing: border-box; }
  body {
    background: #0d1230;
    color: #d6ecf0;
    font-family: "Microsoft YaHei", "微软雅黑", Arial, sans-serif;
    line-height: 1.7;
    padding: 48px 20px;
  }
  .wrap { max-width: 860px; margin: 0 auto; }
  header { text-align: center; margin-bottom: 40px; }
  header h1 { font-size: 2.2em; color: #d6ecf0; letter-spacing: 2px; }
  header h1 span { color: #8a2be2; }
  header p { color: #8a2be2; margin-top: 8px; font-size: 0.95em; }
  .card {
    background: rgba(138, 43, 226, 0.07);
    border: 1px solid rgba(138, 43, 226, 0.45);
    border-radius: 10px;
    padding: 24px 28px;
    margin-bottom: 28px;
  }
  .card h2 { font-size: 1.25em; color: #d6ecf0; border-left: 4px solid #8a2be2; padding-left: 12px; }
  .oneline { margin: 12px 0 6px; }
  .tag { color: #8a2be2; font-family: Consolas, monospace; font-size: 0.9em; margin-bottom: 10px; }
  .card h3 { font-size: 1em; color: #8a2be2; margin: 14px 0 6px; }
  .card ul { list-style: none; }
  .card li { padding-left: 18px; position: relative; margin-bottom: 4px; }
  .card li::before { content: "▸"; color: #8a2be2; position: absolute; left: 0; }
  .debug { padding-left: 18px; border-left: 2px dashed rgba(138, 43, 226, 0.6); }
  .link { margin-top: 12px; font-size: 0.92em; color: #9fb3c8; word-break: break-all; }
  .collab { margin-top: 8px; }
  .collab strong { color: #8a2be2; }
  footer { text-align: center; color: #5a6b8c; font-size: 0.85em; margin-top: 32px; }
  @media print {
    body { background: #fff; color: #222; padding: 20px; }
    .card { border-color: #8a2be2; background: #fff; break-inside: avoid; }
    .card h2, header h1 { color: #222; }
    .link, footer { color: #555; }
  }
</style>
</head>
<body>
  <div class="wrap">
    <header>
      <h1>%s <span>·</span> %s</h1>
      <p>%s</p>
    </header>
%s
    <section class="card">
      <h2>我与 AI 的协作模式</h2>
      <ul class="collab">
%s
      </ul>
    </section>
    <footer>Generated from finalized CV drafts · facts &amp; figures verbatim</footer>
  </div>
</body>
</html>
""" % (esc(PORTFOLIO["name"]), esc(PORTFOLIO["title"]),
       esc(PORTFOLIO["name"]), esc(PORTFOLIO["title"]), esc(PORTFOLIO["contact"]),
       "\n".join(cards), collab)

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    print("saved:", os.path.normpath(out_path))


# ---------------------------------------------------------------- PDF 转换


def convert_pdf(docx_path):
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "pdf",
         "--outdir", OUT_DIR, docx_path],
        check=True, capture_output=True, timeout=180,
    )
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    print("saved:", os.path.normpath(pdf_path))
    return pdf_path


# ---------------------------------------------------------------- 验证


def verify_pdf(pdf_path, max_pages, keywords):
    from pypdf import PdfReader
    reader = PdfReader(pdf_path)
    pages = len(reader.pages)
    text = "".join((pg.extract_text() or "") for pg in reader.pages)
    ok_pages = pages <= max_pages
    missing = [k for k in keywords if k not in text]
    size_kb = os.path.getsize(pdf_path) / 1024
    status = "OK" if ok_pages and not missing else "FAIL"
    print("[%s] %s: %d 页(上限 %d), %.0f KB, 缺失关键词: %s"
          % (status, os.path.basename(pdf_path), pages, max_pages,
             size_kb, missing or "无"))
    return status == "OK"


def verify_docx_fonts(docx_path):
    doc = Document(docx_path)
    total = with_east = 0
    for p in doc.paragraphs:
        for run in p.runs:
            total += 1
            rPr = run._element.rPr
            if rPr is not None and rPr.rFonts is not None and \
                    rPr.rFonts.get(qn("w:eastAsia")) == EAST_FONT:
                with_east += 1
    size_kb = os.path.getsize(docx_path) / 1024
    ok = total > 0 and with_east == total
    print("[%s] %s: %.0f KB, eastAsia=%s 命中 %d/%d run"
          % ("OK" if ok else "FAIL", os.path.basename(docx_path),
             size_kb, EAST_FONT, with_east, total))
    return ok


# ---------------------------------------------------------------- main


def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    docx_paths = []
    for out_name, md_name in CV_FILES.items():
        data = parse_cv(os.path.join(CV_DIR, md_name))
        assert data["name"] and data["skills"] and data["projects"], md_name
        out_path = os.path.join(OUT_DIR, out_name + ".docx")
        build_cv(data, out_path)
        docx_paths.append(out_path)

    portfolio_docx = os.path.join(OUT_DIR, "portfolio.docx")
    build_portfolio_docx(portfolio_docx)
    docx_paths.append(portfolio_docx)
    build_portfolio_html(os.path.join(OUT_DIR, "portfolio.html"))

    print("\n-- PDF 转换 --")
    pdf_paths = [convert_pdf(p) for p in docx_paths]

    print("\n-- 验证 --")
    all_ok = True
    checks = {
        "cv-ai-infra.pdf": (1, ["11,000+", "40,000+"]),
        "cv-iot.pdf": (1, ["11,000+", "40,000+", "85%", "99%"]),
        "cv-ops.pdf": (1, ["11,000+", "40,000+", "85%", "99%"]),
        "portfolio.pdf": (2, ["11,000+", "44,000+", "40,000+"]),
    }
    for pdf_path in pdf_paths:
        max_pages, keywords = checks[os.path.basename(pdf_path)]
        all_ok &= verify_pdf(pdf_path, max_pages, keywords)
    for docx_path in docx_paths:
        all_ok &= verify_docx_fonts(docx_path)

    print("\n全部通过" if all_ok else "\n存在未通过项,请检查")
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
